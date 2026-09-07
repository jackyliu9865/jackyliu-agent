#!/usr/bin/env python3
"""Export a deck's slides as Markdown, one file per slide.

House note, 25 August 2026: "Maybe convert the sample slides into .mds".

A .pptx is a binary. It cannot be read in Obsidian, diffed, searched, linked
from another note, or reviewed without opening PowerPoint. This writes each
slide out as Markdown carrying the argument (title, strapline, exhibit titles,
basis lines, sources, the read) and the layout facts (which master layout, what
sits where in centimetres), so a deck becomes reviewable as text and
indexable in the vault.

    python3 scripts/slides_to_md.py deck.pptx out_dir/ [--index "Name"]

It is a one-way export for review, not a round-trip format. The build script
stays the source of truth for the deck.
"""
import os
import re
import sys

from pptx import Presentation
from pptx.util import Emu

CM = 360000.0
# Chart furniture: values, category names, panel captions. These belong to the
# exhibit, not to the slide's argument, and listing them as body copy turns a
# readable export into a dump.
FURNITURE = {"DUMBBELL_CAT", "DUMBBELL_VALUE", "SHAREBAR_CAP", "PANEL_TITLE",
             "EXHIBIT_KEY", "EXHIBIT_REFLABEL"}

ORDER = ["EXHIBIT_TITLE", "EXHIBIT_BASIS", "EXHIBIT_KEY", "EXHIBIT_EVENT",
         "EXHIBIT_REFLABEL", "EXHIBIT_CALLOUT", "COMMENTARY", "HEADER_BAR",
         "KPI_NUMBER", "KPI_LABEL", "KPI_BASIS", "KEY_FEATURES_HEAD",
         "KEY_FEATURE_LEAD", "KEY_FEATURE_BODY", "EXHIBIT_SOURCE"]


def _cm(v):
    return None if v is None else round(v / CM, 2)


def _slug(text, n=60):
    s = re.sub(r"[^A-Za-z0-9]+", "-", (text or "slide")).strip("-").lower()
    return (s[:n] or "slide")


def _paras(shape):
    out = []
    for p in shape.text_frame.paragraphs:
        t = "".join(r.text for r in p.runs).strip()
        if t:
            out.append((t, p.level))
    return out


def slide_to_md(slide, n, total):
    layout = slide.slide_layout.name
    title = ""
    try:
        if slide.shapes.title is not None:
            title = slide.shapes.title.text_frame.text.strip()
    except AttributeError:
        pass

    named, other, charts, tables = {}, [], [], []
    for sh in slide.shapes:
        nm = sh.name or ""
        if getattr(sh, "has_chart", False):
            ch = sh.chart
            charts.append({
                "id": nm.split("::", 1)[1] if "::" in nm else nm,
                "series": [s.name for s in ch.series],
                "categories": [str(c) for c in list(ch.plots[0].categories)]
                if ch.plots else [],
                "rect": (_cm(sh.left), _cm(sh.top), _cm(sh.width), _cm(sh.height)),
            })
            continue
        if getattr(sh, "has_table", False):
            tables.append([[c.text.strip() for c in row.cells]
                           for row in sh.table.rows])
            continue
        if not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        if sh is getattr(slide.shapes, "title", None):
            continue
        if nm in FURNITURE:
            continue
        key = nm if nm in ORDER else None
        if key:
            named.setdefault(key, []).append(sh)
        else:
            other.append(sh)

    L = []
    L.append("---")
    L.append("tags: [deck-slide, kpmg-deck, exported]")
    L.append("slide: %d" % n)
    L.append("layout: %s" % layout)
    L.append("---")
    L.append("")
    L.append("# %s" % (title or "Slide %d" % n))
    L.append("")
    L.append("> Slide %d of %d, on layout `%s`." % (n, total, layout))
    L.append("")

    def block(key, heading, bullet=True):
        if key not in named:
            return
        L.append("## %s" % heading)
        L.append("")
        for sh in named[key]:
            for t, _lvl in _paras(sh):
                L.append(("- " if bullet else "") + t)
        L.append("")

    # The strapline is the slide's one-line argument. On the master it is
    # placeholder idx 18; where a layout has none, deckkit draws it as a 9pt
    # bold Pacific Blue box in the strapline band, which is how we find it.
    strap, body = None, []
    for sh in other:
        txt = sh.text_frame.text.strip()
        if txt == title:
            continue
        is_strap = False
        try:
            if sh.is_placeholder and sh.placeholder_format.idx == 18:
                is_strap = True
        except (AttributeError, ValueError):
            pass
        try:
            r = sh.text_frame.paragraphs[0].runs[0]
            if (r.font.bold and r.font.size is not None
                    and round(r.font.size.pt) == 9
                    and str(r.font.color.rgb) == "00B8F5"):
                is_strap = True
        except (AttributeError, IndexError, TypeError, ValueError):
            pass
        if is_strap and strap is None:
            strap = txt
        else:
            body.append(sh)

    if strap:
        L.append("**%s**" % strap)
        L.append("")
    if body:
        L.append("## Body")
        L.append("")
        for sh in body:
            for t, lvl in _paras(sh):
                prefix = "- " if lvl >= 2 else ""
                L.append("%s%s%s" % ("  " * max(lvl - 2, 0), prefix, t))
        L.append("")

    block("KPI_NUMBER", "Headline numbers", bullet=True)
    block("KEY_FEATURES_HEAD", "Feature band", bullet=False)
    block("KEY_FEATURE_LEAD", "Feature leads")

    for i, c in enumerate(charts, 1):
        L.append("## Exhibit %d" % i)
        L.append("")
        if "EXHIBIT_TITLE" in named and len(named["EXHIBIT_TITLE"]) >= i:
            L.append("**%s**" % named["EXHIBIT_TITLE"][i - 1].text_frame.text.strip())
            L.append("")
        if "EXHIBIT_BASIS" in named and len(named["EXHIBIT_BASIS"]) >= i:
            L.append("*%s*" % named["EXHIBIT_BASIS"][i - 1].text_frame.text.strip())
            L.append("")
        L.append("- Chart ID: `%s`" % c["id"])
        L.append("- Series: %s" % ", ".join(c["series"]))
        cats = c["categories"]
        if cats:
            shown = [x for x in cats if x][:12]
            L.append("- Categories (%d): %s%s"
                     % (len(cats), ", ".join(shown),
                        " ..." if len(cats) > len(shown) else ""))
        L.append("- Position: %.2f, %.2f cm, %.2f x %.2f cm"
                 % tuple(v or 0 for v in c["rect"]))
        L.append("")

    for i, rows in enumerate(tables, 1):
        L.append("## Table %d" % i)
        L.append("")
        if rows:
            L.append("| " + " | ".join(rows[0]) + " |")
            L.append("|" + "---|" * len(rows[0]))
            for r in rows[1:]:
                L.append("| " + " | ".join(r) + " |")
        L.append("")

    block("COMMENTARY", "The read")
    block("EXHIBIT_CALLOUT", "Callouts")
    block("EXHIBIT_EVENT", "Event markers")
    block("EXHIBIT_SOURCE", "Sources", bullet=False)
    return "\n".join(L).rstrip() + "\n"


def main(pptx_path, out_dir, index_name=None):
    prs = Presentation(pptx_path)
    os.makedirs(out_dir, exist_ok=True)
    slides = list(prs.slides)
    written = []
    for n, slide in enumerate(slides, 1):
        md = slide_to_md(slide, n, len(slides))
        first = ""
        try:
            if slide.shapes.title is not None:
                first = slide.shapes.title.text_frame.text.strip()
        except AttributeError:
            pass
        name = "%02d — %s.md" % (n, (first or slide.slide_layout.name)[:52])
        path = os.path.join(out_dir, name)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(md)
            # Link to the index FILE, not to the deck's display name. Linking
            # to index_name left every exported slide carrying a broken
            # wikilink, because the index is written as "00 — Slide Index.md".
            # Found 25 August 2026 by the vault audit on the HKEX deck.
            fh.write("\n## Related\n\n- [[00 — Slide Index]]\n")
        written.append((n, name, first))

    idx = ["---", "tags: [deck-slide, kpmg-deck, index]",
           "status: active", "---", "",
           "# %s — Slide Index" % (index_name or os.path.basename(pptx_path)),
           "",
           "Exported from `%s` by `slides_to_md.py`. One file per slide, so the "
           "deck can be read, searched and linked without opening PowerPoint."
           % os.path.basename(pptx_path), "",
           "| # | Slide | File |", "|---|---|---|"]
    for n, name, first in written:
        idx.append("| %d | %s | [[%s]] |" % (n, first or "—", name[:-3]))
    idx += ["", "## Related", "", "- [[Vault Index]]", ""]
    with open(os.path.join(out_dir, "00 — Slide Index.md"), "w",
              encoding="utf-8") as fh:
        fh.write("\n".join(idx))
    print("wrote %d slide files to %s" % (len(written), out_dir))
    return written


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    name = None
    if "--index" in sys.argv:
        name = sys.argv[sys.argv.index("--index") + 1]
    main(args[0], args[1], name)
