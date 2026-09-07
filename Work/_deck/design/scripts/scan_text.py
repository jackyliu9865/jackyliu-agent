#!/usr/bin/env python3
"""scan_text.py — dump every string in a deck and flag unfilled placeholders.

Replaces the `markitdown | grep` step in SKILL.md. markitdown needs a separate
install and a Python newer than the macOS system one; python-pptx is already a
hard dependency of deckkit, so this adds nothing to the toolchain.

    python3 scan_text.py out.pptx            # report only
    python3 scan_text.py out.pptx --dump     # also print every slide's text

Exit code 1 if anything suspicious is found, so it can gate a build.
"""

import re
import sys

from pptx import Presentation

# Prompt text left behind when a placeholder is never filled, plus the usual
# drafting debris.
SUSPECT = re.compile(
    r"click to (add|edit)|lorem ipsum|\bTODO\b|\[insert|\bTBC\b|\bXX+\b|placeholder",
    re.I,
)

TITLE_LIMIT = 36  # the 44pt title box holds one line at this width
# Cover and divider titles are 66pt in a much bigger box and take a longer
# budget. Added 26 August 2026: this file applied 36 to every layout and so
# reported both section dividers of a finished deck as too long, while
# deckkit.check() passed them. A gate that contradicts the gate beside it gets
# ignored, which is worse than not having it.
DIVIDER_TITLE_LIMIT = 66

# House style. The deck-relevant subset of the full KPMG house style QRG, which
# lives in the vault at `Career/KPMG House Style — PowerPoint & Report QRG.md` and
# is NOT bundled with this package (it governs every KPMG deliverable, not just
# decks). See kpmg-deck/SKILL.md for the pointer.  (pattern, message)
HOUSE_STYLE = [
    (re.compile(r"\bN/A\b"),                    "write n/a, not N/A"),
    (re.compile(r"\bn\.a\.(?!\w)"),              "write n/a (n.a. means not available)"),
    # "ROIC vs WACC" is the established name of the exhibit and appears in the
    # POV slide plan, so exempt that one phrase rather than flagging every deck
    # built to the plan.
    (re.compile(r"(?<![\w.])vs\.?(?=\s)(?<!ROIC vs)"), "write versus, not vs"),
    (re.compile(r"(?<![\w.])eg(?=[\s,])"),        "write e.g., not eg"),
    (re.compile(r"(?<![\w.])ie(?=[\s,])"),        "write i.e., not ie"),
    (re.compile(r"(?<![\w.])etc(?!\.)"),          "write etc., not etc"),
    (re.compile(r"\b(can't|isn't|doesn't|won't|didn't)\b", re.I), "expand contractions (cannot, is not)"),
    (re.compile(r"\bFY20\d{2}\b"),               "use FY23 form, not FY2023"),
    (re.compile(r"\s+[.,;]"),                    "space before punctuation"),
    (re.compile(r"[a-z0-9]  +[a-zA-Z]"),          "double space"),
    (re.compile(r"[€£$¥]\s+\d"),                  "no space between currency symbol and figure"),
    (re.compile(r"\b(GBP|USD|EUR|HKD|RMB|CNY)\d"), "space between currency code and figure"),
    (re.compile(r"\bCNY\b"),                     "write RMB, not CNY"),
    (re.compile(r"\bslide \d+|\bpage \d+", re.I), "do not reference slide or page numbers; reference the section"),
    # The four QRG find-and-replace rows this list did not cover, added
    # 26 August 2026. Each is written to fire only where the QRG actually
    # objects, because a gate that cries wolf gets ignored and that is worse
    # than not having it. Verified against all three finished decks: zero
    # false positives.
    #
    # "space hyphen space" -> en dash. A hyphen takes no surrounding spaces.
    # An en dash or em dash already in the text is correct and is not matched.
    (re.compile(r"\S \- \S"),                    "use an en dash, not a spaced hyphen"),
    # Ordinals. "1st, 2nd" -> "First, second, third". Single digit only, which
    # matches the golden rule that one to nine go in words and 10 upward in
    # numerals. Corrected 26 August 2026: matching any length flagged "75th
    # percentile" on a finished deck, and "seventy-fifth percentile" would be
    # worse English than the thing it was objecting to.
    (re.compile(r"\b[1-9](?:st|nd|rd|th)\b", re.I), "write ordinals in words: first, second, third"),
    # "circa" is for dates only; anywhere else it is "approximately". A bare
    # "c." before a four-digit year is the sanctioned date use and is allowed,
    # so only "c." before some other number is flagged.
    (re.compile(r"\bcirca\b", re.I),              "write approximately, not circa (circa is for dates only)"),
    (re.compile(r"\bc\.\s*(?!\d{4}\b)\d"),        "write approximately, not c., unless before a year"),
    # "%" as a noun. "gross margin %" is wrong and "60.5%" is right, so this
    # fires only where the sign is loose in prose. Corrected 26 August 2026: a
    # basis line declaring its units as ", %" is the exhibit grammar this skill
    # mandates, and the first version flagged it on nine pages across two
    # finished decks. A comma before the sign means units, not prose.
    # A trailing "(...)" after the sign is the other units-declaration idiom,
    # "USD billion (column, LHS) and % (dot, RHS)", which exhibits.md documents
    # as the sanctioned dual-axis basis line. Added to the exemption 26 August
    # 2026 when the gate flagged that exact worked example.
    (re.compile(r"(?<![,\d\s])\s%(?!\d)(?!\s*\()"),
                                                 "write percentage in text; % only attached to a figure"),
]

# Judgement calls. The QRG asks for these to be *checked*, not corrected on
# sight, so they are reported and do not fail the build. Added 26 August 2026:
# the slash rule sat in the failing list and flagged eleven formulae on a
# finished deck — "NOPAT / Revenue", "Revenue / IC", "Total Revenue /
# Aggregated Working Capital" — where spaces around the slash are how a ratio
# is written and closing them up would be worse. The QRG's own wording is
# "check each instance; sometimes a space or line break is being used to split
# a paragraph", which is a review instruction, not a defect.
HOUSE_STYLE_REVIEW = [
    (re.compile(r"\s/\s"),
     "spaced slash: correct in a ratio or formula, wrong in prose — check it"),
]


# A bullet is body text on a content placeholder. House rule: no terminal full
# stop. Two things are exempt and SHOULD end in a full stop, because both are
# complete sentences rather than list items:
#   idx 17, the source / note strip
#   idx 18, the strapline, which states the slide's argument as an assertion
# (idx 18 added 21 Aug 2026: it was being flagged on every deck that wrote a
# strapline as a sentence, including deckkit's own documented example.)
BULLET_STOP = re.compile(r"[a-zA-Z0-9)\]%]\.$")

# A bullet is one item. A block carrying more than one sentence is prose, and
# prose is punctuated. Added 26 August 2026: the rule was flagging the legal
# disclaimer and copyright block on the back page of a finished deck, which is
# boilerplate that must keep its full stops and is not a bullet at all.
PROSE_BLOCK = re.compile(r"[.!?]\s+[A-Z(\u00a9]")
COPYRIGHT_LINE = re.compile(r"^\s*\u00a9|\ball rights reserved\b|"
                            r"\btrademarks? (?:are )?used under licen[cs]e\b",
                            re.I)


def is_prose_block(text):
    """True where the no-terminal-stop bullet rule does not apply."""
    t = text.strip()
    return bool(PROSE_BLOCK.search(t) or COPYRIGHT_LINE.search(t))
EXEMPT_IDX = {17, 18}

# The strapline band, for the fallback text box deckkit.strapline_box() draws on
# layouts with no idx 18. y 2.72 cm, height 0.9 cm.
STRAPLINE_BAND_EMU = (int(2.4 * 360000), int(3.7 * 360000))


# Anything at or below this y is a source/note strip. Derived from the grid:
# deckkit.source_box() draws the fallback source line at content_bottom - 0.6 =
# 15.72 cm, and placeholder idx 17 sits at 16.32 cm. The old threshold was
# 5,850,000 EMU (16.25 cm), which exempted idx 17 but NOT the fallback box, so
# every hand-drawn source line on `Title only_Blank` failed the no-full-stop
# rule and the script exits 1 to gate the build. Fixed 21 Aug 2026.
NOTE_STRIP_TOP_EMU = 5_580_000  # 15.5 cm


# Exhibit furniture drawn by exhibits.py. These are not bullets and several of
# them are complete sentences that SHOULD end in a full stop: a source line
# under a mid-slide table sits well above the source strip, so position alone
# cannot exempt it. Added 25 Aug 2026 with the data-exhibit engine.
EXEMPT_SHAPE_NAMES = {
    "EXHIBIT_SOURCE",   # a source line, wherever it sits on the page
    "EXHIBIT_BASIS",    # period / units / scope, a fragment not a sentence
    "EXHIBIT_TITLE",    # the finding, sentence case, no terminal stop
    "EXHIBIT_KEY", "EXHIBIT_REFLABEL", "PANEL_TITLE",
    "KPI_NUMBER", "KPI_LABEL", "KPI_BASIS",
    "SHAREBAR_CAP", "DUMBBELL_CAT", "DUMBBELL_VALUE",
}


def is_note_shape(shape):
    """Source/note strips are exempt from the no-full-stop rule. They are either
    placeholder idx 17, a text box drawn in the source strip by
    deckkit.source_box() on a layout that has no idx 17, or a named piece of
    exhibit furniture from exhibits.py."""
    if (getattr(shape, "name", "") or "") in EXEMPT_SHAPE_NAMES:
        return True
    try:
        if shape.placeholder_format.idx in EXEMPT_IDX:
            return True
    except (AttributeError, ValueError):
        pass
    try:
        if shape.top is None:
            return False
        if shape.top >= NOTE_STRIP_TOP_EMU:
            return True
        # deckkit.strapline_box() fallback: a plain text box in the strapline band
        lo, hi = STRAPLINE_BAND_EMU
        return lo <= shape.top <= hi
    except AttributeError:
        return False


def iter_text(slide):
    """Yield (shape, text, level) per paragraph.

    `level` is the outline level of the paragraph, or None for a table cell.
    It carries the bullet-versus-prose distinction: on this master, levels 0
    and 1 are unbulleted paragraphs and level 2 and deeper carry the bullet
    glyph. See the no-full-stop rule below, which needs to tell them apart.
    """
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if text.strip():
                    yield shape, text, para.level
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        yield shape, cell.text, None


def main(path, dump=False):
    prs = Presentation(path)
    findings = []
    reviews = []

    for n, slide in enumerate(prs.slides, start=1):
        layout = slide.slide_layout.name
        if dump:
            print(f"\n--- slide {n} [{layout}] ---")

        seen_idx = set()
        for shape, text, level in iter_text(slide):
            if dump:
                print(f"  {text}")
            if SUSPECT.search(text):
                findings.append(f"slide {n} [{layout}]: unfilled placeholder text: {text[:70]!r}")
            try:
                idx = shape.placeholder_format.idx
                seen_idx.add(idx)
            except (AttributeError, ValueError):
                idx = None

            for pat, msg in HOUSE_STYLE:
                if pat.search(text):
                    findings.append(f"slide {n} [{layout}]: house style, {msg}: {text[:60]!r}")

            for pat, msg in HOUSE_STYLE_REVIEW:
                if pat.search(text):
                    reviews.append(f"slide {n} [{layout}]: {msg}: {text[:60]!r}")

            # No full stops on BULLETS. Body paragraphs are different and DO
            # take one: house instruction, 27 August 2026, refining the rule
            # recorded on 12 August. The QRG's wording has always been
            # "bullets in decks do not take full stops", and a prose paragraph
            # is not a bullet.
            #
            # On this master, levels 0 and 1 are unbulleted paragraphs and
            # level 2 and deeper carry the bullet glyph, so the level is the
            # distinction. `is_prose_block` stays as a second exemption for
            # multi-sentence blocks, which covers table cells and the back
            # page boilerplate, where there is no level to read.
            is_bullet = level is not None and level >= 2
            if (is_bullet and idx not in EXEMPT_IDX and not is_note_shape(shape)
                    and not is_prose_block(text)):
                if BULLET_STOP.search(text.strip()) and len(text.strip()) > 3:
                    findings.append(
                        f"slide {n} [{layout}]: bullet ends in a full stop: {text[:60]!r}")

        title = slide.shapes.title
        if title is not None and title.has_text_frame:
            t = title.text_frame.text.strip()
            low = layout.lower()
            limit = (DIVIDER_TITLE_LIMIT
                     if ("cover" in low or "divider" in low) else TITLE_LIMIT)
            if len(t) > limit:
                findings.append(
                    f"slide {n} [{layout}]: title is {len(t)} chars, limit {limit}: {t[:50]!r}"
                )

        # 18 = strapline, 17 = source line. Content layouts carry both.
        available = {ph.placeholder_format.idx for ph in slide.placeholders}
        if 18 in available and 18 not in seen_idx:
            findings.append(f"slide {n} [{layout}]: empty strapline (idx 18), no stated argument")
        if 17 in available and 17 not in seen_idx:
            findings.append(f"slide {n} [{layout}]: empty source line (idx 17)")

    if reviews:
        print(f"\n{len(reviews)} to review (not failures, the QRG asks you to "
              f"check each one):")
        for r in reviews:
            print(f"  ? {r}")

    if findings:
        print(f"\n{len(findings)} finding(s):")
        for f in findings:
            print(f"  - {f}")
        return 1

    print(f"scan_text: clean ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1], dump="--dump" in sys.argv))
